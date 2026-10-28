"""SignalCraft Streamlit frontend for five workflows and multi-provider AI fallback."""
import json, os
from io import BytesIO
from datetime import date
import requests
import streamlit as st
from dotenv import load_dotenv
load_dotenv()
API=os.getenv("API_URL","https://signalcraft-brand-campaign-studio.onrender.com").rstrip("/")
st.set_page_config(page_title="SignalCraft | Brand & Campaign Studio",page_icon="✨",layout="wide")
st.markdown("""<style>
:root { --sc-ink:#172033; --sc-muted:#64748b; --sc-accent:#6d5dfc; }
.block-container { max-width: 1440px; padding-top: 1.5rem; padding-bottom: 3rem; }
[data-testid="stAppViewContainer"] { background: radial-gradient(circle at 8% 0%, rgba(109,93,252,.10), transparent 28%), radial-gradient(circle at 95% 8%, rgba(14,165,233,.09), transparent 24%); }
[data-testid="stHeader"] { background: transparent; }
h1 { letter-spacing: -0.045em; font-weight: 800 !important; color: var(--sc-ink); }
h2, h3 { letter-spacing: -0.025em; color: var(--sc-ink); }
[data-testid="stCaptionContainer"] p { color: var(--sc-muted); }
[data-testid="stTabs"] [role="tablist"] { gap: .4rem; border-bottom: 1px solid rgba(148,163,184,.25); }
[data-testid="stTabs"] button[role="tab"] { border-radius: 12px 12px 0 0; padding: .7rem 1rem; font-weight: 650; }
[data-testid="stTabs"] button[aria-selected="true"] { background: rgba(109,93,252,.10); }
[data-testid="stVerticalBlockBorderWrapper"] { border-radius: 18px; border-color: rgba(148,163,184,.24); box-shadow: 0 8px 28px rgba(15,23,42,.035); }
[data-testid="stTextArea"] textarea, [data-testid="stTextInput"] input, [data-testid="stNumberInput"] input { border-radius: 12px; }
[data-testid="stFileUploader"] section { border-radius: 14px; border: 1px dashed rgba(109,93,252,.45); background: rgba(109,93,252,.035); }
div.stButton > button { border-radius: 12px; font-weight: 650; transition: transform .12s ease, box-shadow .12s ease; }
div.stButton > button[kind="primary"] { width: 100%; min-height: 2.8rem; border: 0; background: linear-gradient(110deg,#635bff,#8b5cf6); color: white; box-shadow: 0 6px 16px rgba(99,91,255,.20); }
div.stButton > button:hover { transform: translateY(-1px); box-shadow: 0 8px 18px rgba(15,23,42,.10); }
[data-testid="stAlert"] { border-radius: 14px; }
[data-testid="stJson"] { border-radius: 14px; }
footer { visibility: hidden; }
</style>""", unsafe_allow_html=True)

def err(r):
    try:
        d=r.json().get("detail",r.json()); return json.dumps(d,ensure_ascii=False,indent=2) if isinstance(d,(dict,list)) else str(d)
    except Exception:return r.text[:2000]
def get(path):return requests.get(API+path,timeout=25)
def upload_text(f):
    raw=f.getvalue(); n=f.name.lower()
    if n.endswith((".txt",".md",".csv")):
        for enc in ("utf-8-sig","utf-8","cp1252"):
            try:return raw.decode(enc).strip()
            except UnicodeDecodeError:pass
        raise ValueError("Cannot decode file. Save as UTF-8 and retry.")
    if n.endswith(".pdf"):
        from pypdf import PdfReader
        text="\n".join(p.extract_text() or "" for p in PdfReader(BytesIO(raw)).pages).strip()
        if not text:raise ValueError("No selectable text found. Scanned PDF OCR is not included.")
        return text
    if n.endswith(".docx"):
        from docx import Document
        d=Document(BytesIO(raw)); bits=[p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:bits.append(" | ".join(c.text for c in row.cells))
        return "\n".join(bits).strip()
    raise ValueError("Supported file types: TXT, MD, CSV, PDF, DOCX.")
def source_input(key,label,helptext="Paste text or upload a document."):
    text=st.text_area(label,key=key+"_text",height=150,help=helptext)
    f=st.file_uploader("Or upload a source file",type=["txt","md","csv","pdf","docx"],key=key+"_file")
    if f:
        try:
            ft=upload_text(f)
            if text.strip():return text.strip()+"\n\nUploaded source ("+f.name+"):\n"+ft
            return ft
        except Exception as e:st.error(f"File read failed: {e}")
    return text.strip()
def post(path, source, **params):
    try:
        r=requests.post(API+path,params={"source":source,**params},timeout=240)
        if r.ok:return r.json()
        st.error(f"Request failed (HTTP {r.status_code})"); st.code(err(r)); return None
    except requests.RequestException as e:st.error(f"API connection failed: {e}");return None
def show_saved(result, key):
    if result:
        st.success(result.get("message","Saved successfully"))
        if result.get("provider"):st.caption("AI provider used: "+result["provider"])
        st.json(result)
        st.session_state[key]=result

def load_list(path,key):
    try:
        r=get(path)
        if r.ok:
            data=r.json().get(key,[])
            if data:st.json(data)
            else:st.info("Nothing saved yet.")
        else:st.error(err(r))
    except Exception as e:st.error(f"Cannot load saved records: {e}")

st.title("✦ SignalCraft")
st.markdown("**Brand & Campaign Studio**  ·  Turn messy marketing inputs into validated, reusable assets.")
st.caption("MULTI-MODEL AI  /  STRUCTURED OUTPUTS  /  SQLITE WORKSPACE")
try:
    h=get("/health")
    if h.ok:
        x=h.json(); providers=x.get("configured_providers",[])
        st.success("API online · Configured providers: "+(", ".join(providers) if providers else "none — add API keys in .env"))
        if len(providers)<2:st.warning("Fallback works best with at least two configured providers.")
    else:st.warning(f"Backend status {h.status_code}: {err(h)}")
except Exception:st.error("Backend unavailable. Start it with: uvicorn signalcraft:app --reload")

tabs=st.tabs(["📣 Campaign plan","🗓️ Content calendar","🏷️ Brand profile","📊 Performance insights","✍️ Content drafts & versions","🗃️ Saved records"])
with tabs[0]:
    st.subheader("1 · Turn a brief into a structured campaign")
    src=source_input("campaign","Campaign brief", "Describe goals, audience, message, channels, dates and budget. Or upload a TXT/MD/CSV/PDF/DOCX file.")
    if st.button("Generate and save campaign",type="primary",key="campaign_go"):
        if len(src)<10:st.error("Paste a brief or upload a file with at least 10 characters.")
        else:
            with st.spinner("Extracting campaign with provider fallback…"):
                try:
                    r=requests.post(API+"/campaigns",params={"brief":src},timeout=240)
                    if r.ok:
                        d=r.json();st.success(f"Campaign saved · ID {d['id']} · Provider: {d.get('provider','unknown')}");st.json(d["campaign"])
                    else:st.error(f"Campaign failed (HTTP {r.status_code})");st.code(err(r))
                except Exception as e:st.error(str(e))
with tabs[1]:
    st.subheader("2 · Generate a saved content calendar")
    src=source_input("calendar","Marketing objective or campaign brief")
    c1,c2=st.columns(2)
    with c1:count=st.slider("Number of entries",1,30,10)
    with c2:platforms=st.text_input("Platforms", "Instagram,LinkedIn,YouTube,Email")
    if st.button("Generate calendar",type="primary",key="calendar_go"):
        if len(src)<10:st.error("Provide a marketing objective or campaign brief first.")
        else:
            with st.spinner("Creating calendar and saving it…"):show_saved(post("/calendar",src,count=count,platforms=platforms),"last_calendar")
with tabs[2]:
    st.subheader("3 · Extract a reusable brand profile")
    src=source_input("brand","Brand document or brand notes", "Upload a brand guide, product description, or paste notes about brand voice and audience.")
    if st.button("Extract and save brand profile",type="primary",key="brand_go"):
        if len(src)<10:st.error("Provide a brand document or notes first.")
        else:
            with st.spinner("Extracting brand profile…"):show_saved(post("/brand-profiles",src),"last_brand")
with tabs[3]:
    st.subheader("4 · Turn performance notes into insights and actions")
    src=source_input("performance","Campaign performance notes", "Paste campaign results, KPI notes, or a report. The AI will not invent missing metrics.")
    if st.button("Analyze and save insights",type="primary",key="performance_go"):
        if len(src)<10:st.error("Provide campaign performance notes first.")
        else:
            with st.spinner("Analyzing metrics and saving insights…"):show_saved(post("/performance-insights",src),"last_performance")
with tabs[4]:
    st.subheader("5 · Create platform drafts and keep review versions")
    src=source_input("drafts","Approved campaign brief and relevant brand voice", "Include the campaign facts and any brand profile details you want the drafts to follow.")
    platforms=st.text_input("Target platforms", "Instagram,LinkedIn,YouTube",key="draft_platforms")
    if st.button("Generate platform drafts",type="primary",key="drafts_go"):
        if len(src)<10:st.error("Provide the approved brief first.")
        else:
            with st.spinner("Generating platform-specific drafts…"):show_saved(post("/content-drafts",src,platforms=platforms),"last_drafts")
    st.markdown("#### Save a revised version")
    draft_id=st.number_input("Saved draft bundle ID",min_value=1,step=1,value=1,key="draft_id")
    edited=st.text_area("Revised draft bundle JSON",height=180,placeholder='Paste the complete edited draft_bundle JSON object here.')
    note=st.text_input("Version note","Copy edits after review")
    if st.button("Save as new version",key="save_version"):
        try:
            obj=json.loads(edited)
            r=requests.put(f"{API}/content-drafts/{int(draft_id)}",json={"content":obj,"change_note":note},timeout=30)
            if r.ok:st.success(f"Saved version {r.json().get('version')}")
            else:st.error(err(r))
        except Exception as e:st.error(f"Could not save version: {e}")
    if st.button("View draft versions",key="view_versions"):
        try:
            r=get(f"/content-drafts/{int(draft_id)}/versions")
            if r.ok:st.json(r.json())
            else:st.error(err(r))
        except Exception as e:st.error(str(e))
with tabs[5]:
    st.subheader("Saved records")
    kinds={"Campaigns":"/campaigns","Content calendars":"/calendar","Brand profiles":"/brand-profiles","Performance reports":"/performance-insights","Content draft bundles":"/content-drafts"}
    choice=st.selectbox("Choose record type",list(kinds))
    if st.button("Load saved records",type="primary"):load_list(kinds[choice],{"Campaigns":"campaigns","Content calendars":"plans","Brand profiles":"profiles","Performance reports":"reports","Content draft bundles":"drafts"}[choice])
